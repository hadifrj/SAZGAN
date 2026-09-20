"""
Sazgan V2.4 Excel Engine
Generic Excel import pipeline. Final column mappings can be configured when the user's
final Excel workbook is supplied.
"""
from __future__ import annotations
import json, re, sqlite3, hashlib
from pathlib import Path
from datetime import datetime, timezone
from typing import Any, Dict, List

try:
    import openpyxl
except ImportError:
    openpyxl = None

def _now(): return datetime.now(timezone.utc).isoformat()

def normalize_text(v):
    if v is None: return ""
    s=str(v).replace("\u200c"," ").replace("\u200f","").replace("\u200e","")
    s=re.sub(r"\s+"," ",s).strip()
    return s

def normalize_header(v):
    s=normalize_text(v).lower()
    return re.sub(r"[\s_\-]+","",s)

def load_workbook(path):
    if openpyxl is None:
        raise RuntimeError("openpyxl is required for Excel import")
    return openpyxl.load_workbook(path, read_only=True, data_only=True)

def read_sheet(path, sheet=None, max_rows=None):
    wb=load_workbook(path)
    ws=wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows=list(ws.iter_rows(values_only=True))
    if not rows: return {"sheet":ws.title,"headers":[],"rows":[]}
    headers=[normalize_text(x) for x in rows[0]]
    data=[]
    for raw in rows[1:max_rows+1 if max_rows else None]:
        if all(normalize_text(x)=="" for x in raw): continue
        item={}
        for i,h in enumerate(headers):
            if h:
                item[h]=normalize_text(raw[i] if i<len(raw) else "")
        data.append(item)
    return {"sheet":ws.title,"headers":headers,"rows":data}

def fingerprint(row, keys):
    raw="|".join(normalize_text(row.get(k,"")) for k in keys)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()

def validate_rows(rows, required=(), unique_keys=()):
    errors=[]; seen={}
    for idx,row in enumerate(rows, start=2):
        missing=[k for k in required if normalize_text(row.get(k,""))==""]
        if missing:
            errors.append({"row":idx,"type":"required","fields":missing})
        if unique_keys:
            fp=fingerprint(row,unique_keys)
            if fp in seen:
                errors.append({"row":idx,"type":"duplicate","duplicate_of":seen[fp],"keys":list(unique_keys)})
            else: seen[fp]=idx
    return errors

def preview(path, sheet=None, required=(), unique_keys=(), limit=100):
    result=read_sheet(path,sheet)
    rows=result["rows"]
    errors=validate_rows(rows,required,unique_keys)
    return {"ok":not errors,"sheet":result["sheet"],"headers":result["headers"],
            "total_rows":len(rows),"preview":rows[:limit],"errors":errors[:500]}

def mapping_suggestions(headers, target_fields):
    out={}
    norm={normalize_header(h):h for h in headers}
    for field in target_fields:
        nf=normalize_header(field)
        if nf in norm: out[field]=norm[nf]
        else:
            for h in headers:
                nh=normalize_header(h)
                if nf and (nf in nh or nh in nf):
                    out[field]=h; break
    return out
