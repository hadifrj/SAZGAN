# -*- coding: utf-8 -*-
"""Incremental Excel importer for Sazgan.

Designed for a daily export from another system:
- customer key: کد طرف حساب جدید
- service key: شماره پذیرش
- New / Changed / Unchanged / Review classification
- no duplicate services on repeated imports
- preview before commit
- import batch history
"""
from __future__ import annotations
import hashlib, json, os, re, shutil, sqlite3, uuid
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

from core.constants import BASE_DIR, BACKUP_DIR

IMPORT_DIR = Path(BASE_DIR) / "excel_imports"
IMPORT_DIR.mkdir(parents=True, exist_ok=True)

SERVICE_SHEET = "همه سرویس‌ها"
CUSTOMER_SHEET = "لیست مشتری"

SERVICE_HEADERS = [
    "شماره پذیرش","نام مشتری","کد طرف حساب جدید","وضعیت تطبیق","نیاز به بازبینی",
    "ادرس","سریال","سریال_اصلی","وضعیت اصلاح سریال","رایگان / غیر رایگان",
    "نوع سرویس","شهر","تاریخ نصب","استان","محل سرویس","وضعیا","تاریخ پذیرش",
    "تاریخ گزارش فنی","تاریخ مراجعه / پایان","شماره فاکتور","مدل","وسایل همراه",
    "درخواست مشتری","شرح کنترل کیفیت","شرح گزارش تکنسین","شرح اقدامات انجام شده"
]

def _norm(v):
    if v is None: return ""
    s = str(v).strip()
    return s.replace("ي","ی").replace("ك","ک").replace("\u200c"," ")

def _key(v):
    s = _norm(v)
    if not s: return ""
    return re.sub(r"\s+", "", s)

def _fingerprint(row: dict) -> str:
    raw = json.dumps(row, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()

def _backup_db(db_path: str) -> str | None:
    if not os.path.exists(db_path):
        return None
    target_dir = Path(BACKUP_DIR) / "excel_import"
    target_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = target_dir / f"before_excel_import_{stamp}.db"
    src = sqlite3.connect(db_path)
    dst = sqlite3.connect(str(target))
    try:
        src.backup(dst)
    finally:
        dst.close(); src.close()
    return str(target)

def _headers(ws):
    return [_norm(x) for x in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

def _rows(ws):
    headers = _headers(ws)
    for idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        yield idx, {headers[i] if i < len(headers) else f"col_{i+1}": values[i] for i in range(len(values))}

def _find(headers, candidates):
    normalized = {_key(h): h for h in headers}
    for c in candidates:
        k = _key(c)
        if k in normalized: return normalized[k]
    return None

def _customer_rows(wb):
    if CUSTOMER_SHEET not in wb.sheetnames:
        return []
    ws = wb[CUSTOMER_SHEET]
    headers = _headers(ws)
    code_h = _find(headers, ["کد طرف حساب جدید","کد طرف حساب","کد"])
    name_h = _find(headers, ["نام مشتری","نام"])
    province_h = _find(headers, ["استان"])
    city_h = _find(headers, ["شهر"])
    address_h = _find(headers, ["آدرس","ادرس"])
    phone_h = _find(headers, ["تلفن ثابت","تلفن"])
    mobile_h = _find(headers, ["تلفن همراه","موبایل"])
    person_h = _find(headers, ["نوع"])
    out=[]
    for r, row in _rows(ws):
        code=_norm(row.get(code_h)) if code_h else ""
        if not code: continue
        out.append({
            "source_row": r, "external_customer_code": code,
            "name": _norm(row.get(name_h)) if name_h else "",
            "province": _norm(row.get(province_h)) if province_h else "",
            "city": _norm(row.get(city_h)) if city_h else "",
            "address": _norm(row.get(address_h)) if address_h else "",
            "phone": _norm(row.get(phone_h)) if phone_h else "",
            "mobile": _norm(row.get(mobile_h)) if mobile_h else "",
            "person_type": _norm(row.get(person_h)) if person_h else "",
        })
    return out

def _service_rows(wb):
    if SERVICE_SHEET not in wb.sheetnames:
        raise ValueError(f"شیت «{SERVICE_SHEET}» پیدا نشد.")
    ws=wb[SERVICE_SHEET]
    headers=_headers(ws)
    reception_h=_find(headers, ["شماره پذیرش"])
    customer_code_h=_find(headers, ["کد طرف حساب جدید","کد طرف حساب"])
    customer_h=_find(headers, ["نام مشتری"])
    serial_h=_find(headers, ["سریال"])
    serial_orig_h=_find(headers, ["سریال_اصلی","سریال اصلی"])
    type_h=_find(headers, ["نوع سرویس"])
    status_h=_find(headers, ["وضعیا","وضعیت"])
    review_h=_find(headers, ["نیاز به بازبینی"])
    out=[]
    for r,row in _rows(ws):
        reception=_norm(row.get(reception_h)) if reception_h else ""
        if not reception: continue
        raw={_norm(k): (v.isoformat() if hasattr(v,"isoformat") else v) for k,v in row.items()}
        out.append({
            "source_row": r, "external_reception_no": reception,
            "external_customer_code": _norm(row.get(customer_code_h)) if customer_code_h else "",
            "customer_name": _norm(row.get(customer_h)) if customer_h else "",
            "serial": _norm(row.get(serial_h)) if serial_h else "",
            "serial_original": _norm(row.get(serial_orig_h)) if serial_orig_h else "",
            "service_type": _norm(row.get(type_h)) if type_h else "",
            "status_raw": _norm(row.get(status_h)) if status_h else "",
            "needs_review_raw": _norm(row.get(review_h)) if review_h else "",
            "raw": raw,
        })
    return out

def analyze(path: str) -> dict:
    if openpyxl is None:
        raise RuntimeError("openpyxl نصب نیست.")
    wb=openpyxl.load_workbook(path, data_only=True, read_only=True)
    if CUSTOMER_SHEET not in wb.sheetnames or SERVICE_SHEET not in wb.sheetnames:
        missing=[x for x in (CUSTOMER_SHEET,SERVICE_SHEET) if x not in wb.sheetnames]
        raise ValueError("شیت‌های لازم پیدا نشدند: " + "، ".join(missing))
    customers=_customer_rows(wb)
    services=_service_rows(wb)
    conn=sqlite3.connect(os.path.join(BASE_DIR,"sazgan.db"))
    conn.row_factory=sqlite3.Row
    try:
        customer_map={_key(r["external_code"]): r for r in conn.execute(
            "SELECT id, external_code, name, phone, province, city, address, person_type FROM customers WHERE external_code IS NOT NULL AND TRIM(external_code)<>''")}
        reception_map={_key(r["external_reception_no"]): r for r in conn.execute(
            "SELECT id, external_reception_no, customer_id, serial_number, problem_desc, status, device_type, device_model, customer_name, customer_address, customer_phone, province, city, service_type, reception_date FROM requests WHERE external_reception_no IS NOT NULL AND TRIM(external_reception_no)<>''")}
        c_counts={"new":0,"changed":0,"unchanged":0,"review":0}
        c_items=[]
        for c in customers:
            k=_key(c["external_customer_code"]); old=customer_map.get(k)
            if not old: state="new"
            else:
                changed=any(_norm(c.get(f)) and _norm(c.get(f)) != _norm(old[f] or "") for f in ["name","phone","province","city","address","person_type"])
                state="changed" if changed else "unchanged"
            c_counts[state]+=1
            c_items.append({**c,"state":state})
        s_counts={"new":0,"changed":0,"unchanged":0,"review":0}
        s_items=[]
        for s in services:
            k=_key(s["external_reception_no"]); old=reception_map.get(k)
            review=bool(s["needs_review_raw"]) or not s["external_customer_code"]
            if review:
                state="review"
            elif not old:
                state="new"
            else:
                newfp=_fingerprint({k:v for k,v in s.items() if k!="raw"})
                oldfp=_fingerprint({
                    "external_reception_no": old["external_reception_no"],
                    "external_customer_code": next((c["external_customer_code"] for c in customers if c["external_customer_code"] and old["customer_id"] and False), ""),
                    "customer_name": old["customer_name"], "serial": old["serial_number"],
                    "service_type": old["service_type"], "status_raw": old["status"],
                })
                # Compare the important business fields directly.
                changed = (
                    _norm(s["customer_name"]) != _norm(old["customer_name"]) or
                    _norm(s["serial"]) != _norm(old["serial_number"]) or
                    _norm(s["service_type"]) != _norm(old["service_type"]) or
                    _norm(s["status_raw"]) != _norm(old["status"]) or
                    _norm(s["raw"].get("ادرس")) != _norm(old["customer_address"]) or
                    _norm(s["raw"].get("تاریخ پذیرش")) != _norm(old["reception_date"])
                )
                state="changed" if changed else "unchanged"
            s_counts[state]+=1
            s_items.append({**s,"state":state})
        return {
            "file": os.path.basename(path), "customers": c_counts,
            "services": s_counts, "customer_rows": c_items, "service_rows": s_items,
            "total_customers": len(customers), "total_services": len(services),
        }
    finally:
        conn.close()

def _ensure_columns(conn):
    # Migration 0003 should already do this, but keep the importer defensive.
    for table,col in [("customers","external_code"),("requests","external_reception_no")]:
        cols={r[1] for r in conn.execute(f"PRAGMA table_info({table})")}
        if col not in cols:
            conn.execute(f"ALTER TABLE {table} ADD COLUMN {col} TEXT")
    conn.execute("""CREATE TABLE IF NOT EXISTS import_batches (
        id INTEGER PRIMARY KEY AUTOINCREMENT, source_file TEXT NOT NULL,
        file_sha256 TEXT NOT NULL, imported_at TEXT NOT NULL,
        customers_new INTEGER DEFAULT 0, customers_changed INTEGER DEFAULT 0,
        customers_unchanged INTEGER DEFAULT 0, services_new INTEGER DEFAULT 0,
        services_changed INTEGER DEFAULT 0, services_unchanged INTEGER DEFAULT 0,
        services_review INTEGER DEFAULT 0, status TEXT NOT NULL DEFAULT 'completed',
        backup_path TEXT
    )""")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_import_batch_hash ON import_batches(file_sha256)")
    conn.commit()

def commit(path: str) -> dict:
    analysis=analyze(path)
    with open(path,"rb") as f: sha=hashlib.sha256(f.read()).hexdigest()
    conn=sqlite3.connect(os.path.join(BASE_DIR,"sazgan.db"))
    conn.row_factory=sqlite3.Row
    try:
        _ensure_columns(conn)
        previous=conn.execute("SELECT id FROM import_batches WHERE file_sha256=?",(sha,)).fetchone()
        if previous:
            return {"status":"duplicate_file","batch_id":previous["id"], **analysis}
        backup=_backup_db(os.path.join(BASE_DIR,"sazgan.db"))
        cur=conn.cursor()
        counts={k:v for k,v in analysis["customers"].items()}
        scounts={k:v for k,v in analysis["services"].items()}
        cur.execute("""INSERT INTO import_batches(source_file,file_sha256,imported_at,
            customers_new,customers_changed,customers_unchanged,services_new,services_changed,
            services_unchanged,services_review,status,backup_path)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (os.path.basename(path),sha,datetime.now().isoformat(timespec="seconds"),
             counts["new"],counts["changed"],counts["unchanged"],scounts["new"],scounts["changed"],
             scounts["unchanged"],scounts["review"],"running",backup))
        batch_id=cur.lastrowid

        customer_ids={}
        for c in analysis["customer_rows"]:
            k=_key(c["external_customer_code"])
            old=cur.execute("SELECT id FROM customers WHERE external_code=?",(c["external_customer_code"],)).fetchone()
            if c["state"]=="new":
                cur.execute("""INSERT INTO customers(name,phone,province,city,address,
                    equipment_manager_mobile,person_type,external_code)
                    VALUES(?,?,?,?,?,?,?,?)""",
                    (c["name"],c["phone"],c["province"],c["city"],c["address"],c["mobile"] or None,c["person_type"] or "حقیقی",c["external_customer_code"]))
                customer_ids[k]=cur.lastrowid
            elif old:
                customer_ids[k]=old["id"]
                if c["state"]=="changed":
                    cur.execute("""UPDATE customers SET name=?,phone=?,province=?,city=?,address=?,
                        equipment_manager_mobile=?,person_type=? WHERE id=?""",
                        (c["name"],c["phone"],c["province"],c["city"],c["address"],c["mobile"] or None,c["person_type"] or "حقیقی",old["id"]))
        # refresh all customer ids
        for c in analysis["customer_rows"]:
            k=_key(c["external_customer_code"])
            if k and k not in customer_ids:
                row=cur.execute("SELECT id FROM customers WHERE external_code=?",(c["external_customer_code"],)).fetchone()
                if row: customer_ids[k]=row["id"]

        for s in analysis["service_rows"]:
            cust_id=customer_ids.get(_key(s["external_customer_code"]))
            old=cur.execute("SELECT * FROM requests WHERE external_reception_no=?",(s["external_reception_no"],)).fetchone()
            raw=s["raw"]
            address=_norm(raw.get("ادرس"))
            device_type=_norm(s["service_type"]) or "نامشخص"
            # Map the source into existing request fields; preserve every source column in raw JSON.
            vals=(s["customer_name"],address,None,device_type,s["serial"],_norm(raw.get("درخواست مشتری")) or "Import Excel",
                  _norm(raw.get("محل سرویس")) or "در محل",_norm(raw.get("استان")),s["status_raw"] or "جدید",
                  _norm(raw.get("تاریخ پذیرش")), _norm(raw.get("شهر")), _norm(raw.get("مدل")),
                  _norm(raw.get("شماره فاکتور")), _norm(raw.get("وسایل همراه")), _norm(raw.get("شرح کنترل کیفیت")),
                  _norm(raw.get("شرح گزارش تکنسین")), _norm(raw.get("شرح اقدامات انجام شده")),cust_id)
            if not old:
                cur.execute("""INSERT INTO requests(customer_name,customer_address,customer_phone,device_type,
                    serial_number,problem_desc,service_type,province,status,reception_date,city,device_model,
                    invoice_number,accompanying_items,final_qc_desc,technician_desc,actions_taken,customer_id,
                    external_reception_no)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", vals + (s["external_reception_no"],))
                req_id=cur.lastrowid
            elif s["state"]=="changed":
                cur.execute("""UPDATE requests SET customer_name=?,customer_address=?,device_type=?,
                    serial_number=?,problem_desc=?,service_type=?,province=?,status=?,reception_date=?,
                    city=?,device_model=?,invoice_number=?,accompanying_items=?,final_qc_desc=?,
                    technician_desc=?,actions_taken=?,customer_id=? WHERE id=?""",
                    vals[0:1]+vals[1:2]+vals[3:18]+(old["id"],))
                req_id=old["id"]
            else:
                req_id=old["id"]
            cur.execute("""INSERT INTO service_import_records(
                source_file,source_sheet,source_row,external_customer_code,external_reception_no,
                customer_id,request_id,match_status,needs_review,serial_original,serial_normalized,
                service_type_raw,raw_json) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (os.path.basename(path),SERVICE_SHEET,s["source_row"],s["external_customer_code"],
                 s["external_reception_no"],cust_id,req_id,s["state"],1 if s["state"]=="review" else 0,
                 s["serial_original"],s["serial"],s["service_type"],json.dumps(s["raw"],ensure_ascii=False,default=str)))
        cur.execute("UPDATE import_batches SET status='completed' WHERE id=?",(batch_id,))
        conn.commit()
        return {"status":"completed","batch_id":batch_id, "backup":backup, **analysis}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

def save_upload(file_storage) -> str:
    token=uuid.uuid4().hex
    path=IMPORT_DIR / f"{token}.xlsx"
    file_storage.save(path)
    return str(path)

def cleanup_upload(path: str):
    try: Path(path).unlink(missing_ok=True)
    except Exception: pass
