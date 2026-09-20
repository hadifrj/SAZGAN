# -*- coding: utf-8 -*-
"""مدیریت لیست جغرافیایی: استان / شهرستان / شهر

جداول:
  geo_provinces (id, name)
  geo_counties  (id, province_id, name)
  geo_cities    (id, province_id, county_id, name)

Import از اکسل با ستون‌های: استان | شهرستان | شهر
"""
from __future__ import annotations

from core.constants import PROVINCES
from core.db import _cache_get, _cache_set, _cache_invalidate

# داده‌های جغرافیایی به‌ندرت تغییر می‌کنن ولی توابع زیر تقریباً در هر
# رندر صفحه‌ای که فیلد استان/شهرستان/شهر داره صدا زده می‌شن (فرم‌های
# سرویس، تنظیمات، مشتریان، نمایندگان، ...)؛ کش کوتاه‌مدت از همون
# مکانیزم TTL که برای get_setting استفاده می‌شه، تعداد کوئری‌های
# تکراری روی این جداول رو به‌شدت کم می‌کنه.
_GEO_CACHE_PREFIX = 'geo::'


def ensure_geo_schema(conn):
    """ایجاد جداول جغرافیایی + ستون county روی جداول اصلی."""
    conn.execute('''
        CREATE TABLE IF NOT EXISTS geo_provinces (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS geo_counties (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            province_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            UNIQUE(province_id, name),
            FOREIGN KEY (province_id) REFERENCES geo_provinces(id)
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS geo_cities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            province_id INTEGER NOT NULL,
            county_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            UNIQUE(county_id, name),
            FOREIGN KEY (province_id) REFERENCES geo_provinces(id),
            FOREIGN KEY (county_id) REFERENCES geo_counties(id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_geo_counties_prov ON geo_counties(province_id)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_geo_cities_county ON geo_cities(county_id)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_geo_cities_prov ON geo_cities(province_id)')

    try:
        from core.distance import ensure_distance_schema
        ensure_distance_schema(conn)
    except Exception:
        pass

    # Fixed table whitelist; these identifiers are internal schema names, not user input.
    # ستون شهرستان روی جداول اصلی
    for table in ('requests', 'customers', 'representatives'):
        try:
            cols = {r[1] for r in conn.execute(f'PRAGMA table_info({table})').fetchall()}
            if 'county' not in cols:
                conn.execute(f'ALTER TABLE {table} ADD COLUMN county TEXT')
        except Exception:
            pass
    conn.commit()


def geo_has_data(conn) -> bool:
    try:
        row = conn.execute('SELECT COUNT(*) AS c FROM geo_provinces').fetchone()
        return bool(row and row['c'] > 0)
    except Exception:
        return False


def get_geo_provinces(conn=None):
    """لیست نام استان‌ها — از جدول geo یا ثابت PROVINCES (کش‌شده)."""
    cache_key = _GEO_CACHE_PREFIX + 'provinces'
    cached = _cache_get(cache_key)
    if cached is not None:
        return cached
    try:
        if conn is not None and geo_has_data(conn):
            rows = conn.execute('SELECT name FROM geo_provinces ORDER BY name').fetchall()
            result = [r['name'] for r in rows]
            _cache_set(cache_key, result)
            return result
    except Exception:
        pass
    result = list(PROVINCES)
    _cache_set(cache_key, result)
    return result


def get_geo_counties(conn, province_name: str):
    if not province_name:
        return []
    cache_key = _GEO_CACHE_PREFIX + 'counties::' + province_name
    cached = _cache_get(cache_key)
    if cached is not None:
        return cached
    rows = conn.execute(
        '''SELECT c.name FROM geo_counties c
           JOIN geo_provinces p ON p.id = c.province_id
           WHERE p.name = ?
           ORDER BY c.name''',
        (province_name,),
    ).fetchall()
    result = [r['name'] for r in rows]
    _cache_set(cache_key, result)
    return result


def get_geo_cities(conn, province_name: str, county_name: str = ''):
    if not province_name:
        return []
    cache_key = _GEO_CACHE_PREFIX + 'cities::' + province_name + '::' + county_name
    cached = _cache_get(cache_key)
    if cached is not None:
        return cached
    if county_name:
        rows = conn.execute(
            '''SELECT ci.name FROM geo_cities ci
               JOIN geo_provinces p ON p.id = ci.province_id
               JOIN geo_counties c ON c.id = ci.county_id
               WHERE p.name = ? AND c.name = ?
               ORDER BY ci.name''',
            (province_name, county_name),
        ).fetchall()
    else:
        rows = conn.execute(
            '''SELECT ci.name FROM geo_cities ci
               JOIN geo_provinces p ON p.id = ci.province_id
               WHERE p.name = ?
               ORDER BY ci.name''',
            (province_name,),
        ).fetchall()
    result = [r['name'] for r in rows]
    _cache_set(cache_key, result)
    return result


def geo_stats(conn):
    try:
        p = conn.execute('SELECT COUNT(*) AS c FROM geo_provinces').fetchone()['c']
        c = conn.execute('SELECT COUNT(*) AS c FROM geo_counties').fetchone()['c']
        ci = conn.execute('SELECT COUNT(*) AS c FROM geo_cities').fetchone()['c']
        return {'provinces': p, 'counties': c, 'cities': ci}
    except Exception:
        return {'provinces': 0, 'counties': 0, 'cities': 0}


def import_geo_from_rows(conn, rows, clear_existing=True):
    """
    Import از لیست dict با کلیدهای: استان / شهرستان / شهر
    (یا province / county / city)
    """
    def _val(row, *keys):
        for k in keys:
            v = row.get(k)
            if v is not None and str(v).strip():
                return str(v).strip()
        return ''

    parsed = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        # normalize keys
        norm = {}
        for k, v in row.items():
            if k is None:
                continue
            norm[str(k).strip()] = v
        province = _val(norm, 'استان', 'province', 'Province')
        county = _val(norm, 'شهرستان', 'county', 'County', 'شهرستان ')
        city = _val(norm, 'شهر', 'city', 'City')
        if not province:
            continue
        parsed.append((province, county, city))

    if not parsed:
        return 0, 0, 0, 'هیچ ردیف معتبری پیدا نشد'

    if clear_existing:
        conn.execute('DELETE FROM geo_cities')
        conn.execute('DELETE FROM geo_counties')
        conn.execute('DELETE FROM geo_provinces')

    prov_cache = {}
    county_cache = {}  # (province_id, name) -> id
    city_count = 0
    county_count = 0
    prov_count = 0

    for province, county, city in parsed:
        if province not in prov_cache:
            existing = conn.execute(
                'SELECT id FROM geo_provinces WHERE name = ?', (province,)
            ).fetchone()
            if existing:
                prov_cache[province] = existing['id']
            else:
                cur = conn.execute(
                    'INSERT INTO geo_provinces (name) VALUES (?)', (province,)
                )
                prov_cache[province] = cur.lastrowid
                prov_count += 1
        pid = prov_cache[province]

        if county:
            ckey = (pid, county)
            if ckey not in county_cache:
                existing = conn.execute(
                    'SELECT id FROM geo_counties WHERE province_id = ? AND name = ?',
                    (pid, county),
                ).fetchone()
                if existing:
                    county_cache[ckey] = existing['id']
                else:
                    cur = conn.execute(
                        'INSERT INTO geo_counties (province_id, name) VALUES (?, ?)',
                        (pid, county),
                    )
                    county_cache[ckey] = cur.lastrowid
                    county_count += 1
            cid = county_cache[ckey]

            if city:
                existing = conn.execute(
                    'SELECT id FROM geo_cities WHERE county_id = ? AND name = ?',
                    (cid, city),
                ).fetchone()
                if not existing:
                    conn.execute(
                        'INSERT INTO geo_cities (province_id, county_id, name) VALUES (?, ?, ?)',
                        (pid, cid, city),
                    )
                    city_count += 1

    conn.commit()
    _cache_invalidate()  # کش استان/شهرستان/شهر بعد از ایمپورت جدید معتبر نیست
    return prov_count, county_count, city_count, None


def import_geo_from_excel(conn, file_storage, clear_existing=True):
    """Import از فایل اکسل آپلودشده."""
    try:
        import openpyxl
    except ImportError:
        return 0, 0, 0, 'کتابخانه openpyxl نصب نیست'

    try:
        wb = openpyxl.load_workbook(file_storage, read_only=True, data_only=True)
    except Exception as e:
        return 0, 0, 0, f'خطا در خواندن فایل: {e}'

    # Prefer sheet with استان/شهرستان/شهر
    ws = None
    for name in wb.sheetnames:
        if 'کامل' in name or 'شهر' in name or 'all' in name.lower():
            ws = wb[name]
            break
    if ws is None:
        ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return 0, 0, 0, 'فایل خالی است'

    # Skip title row if first cell looks like a title
    header_list = [str(h).strip() if h is not None else '' for h in header]
    if not any(h in ('استان', 'province', 'Province', 'شهرستان', 'شهر') for h in header_list):
        try:
            header = next(rows_iter)
            header_list = [str(h).strip() if h is not None else '' for h in header]
        except StopIteration:
            return 0, 0, 0, 'هدر معتبر پیدا نشد'

    # Map columns
    col_map = {}
    for i, h in enumerate(header_list):
        hl = h.lower()
        if h in ('استان',) or hl == 'province':
            col_map['استان'] = i
        elif h in ('شهرستان',) or hl == 'county':
            col_map['شهرستان'] = i
        elif h in ('شهر',) or hl == 'city':
            col_map['شهر'] = i

    if 'استان' not in col_map:
        return 0, 0, 0, 'ستون «استان» در فایل پیدا نشد'

    rows = []
    for row in rows_iter:
        if not row:
            continue
        item = {}
        for key, idx in col_map.items():
            if idx < len(row):
                item[key] = row[idx]
        if item.get('استان'):
            rows.append(item)

    return import_geo_from_rows(conn, rows, clear_existing=clear_existing)
