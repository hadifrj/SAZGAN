# -*- coding: utf-8 -*-
"""مسیرهای دستمزد و قطعات: wages(حذف)، parts(حذف/ورودی/خروجی)."""
# -*- coding: utf-8 -*-
from __future__ import annotations

import logging

from flask import (
    request, session, redirect, render_template, jsonify, send_file, current_app
)

logger = logging.getLogger("sazgan")
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from io import BytesIO
import os, re, sqlite3, base64, time
import jdatetime
try:
    import openpyxl
except ImportError:
    openpyxl = None

from core.constants import *
from core import helpers as H


from core.helpers import *  # noqa: F401,F403
from core.constants import *  # noqa: F401,F403



def register(app):
    @app.route('/settings/wages', methods=['GET', 'POST'])
    def settings_wages():
        current_user = get_current_user()
        if not current_user:
            return redirect('/login')
        if not (_is_reception(current_user) or user_has_access(current_user, 'parties') or user_has_access(current_user, 'wages_settings')):
            return redirect('/')

        from core.wages import ensure_wage_schema, seed_default_wage_rates
        conn = get_db()
        ensure_wage_schema(conn)
        msg = error = None

        if request.method == 'POST':
            action = (request.form.get('action') or 'add').strip()
            try:
                if action == 'seed_defaults':
                    seed_default_wage_rates(conn, force=False)
                    msg = 'جدول پیش‌فرض بارگذاری / به‌روز شد'
                elif action == 'edit':
                    wid = int(request.form.get('wage_id') or 0)
                    outside_raw = request.form.get('amount_outside')
                    outside = int(outside_raw) if outside_raw not in (None, '') else None
                    conn.execute(
                        'UPDATE wage_rates SET work_description=?, wage_amount=?, rate_code=?, amount_outside=?, is_per_km=?, category=?, sort_order=?, device_type=?, install_slot=?, is_central_device=? WHERE id=?',
                        (
                            request.form.get('work_description'),
                            int(request.form.get('wage_amount') or 0),
                            request.form.get('rate_code') or None,
                            outside,
                            1 if request.form.get('is_per_km') else 0,
                            request.form.get('category') or None,
                            int(request.form.get('sort_order') or 0),
                            (request.form.get('device_type') or '').strip() or None,
                            (request.form.get('install_slot') or '').strip() or None,
                            1 if (request.form.get('install_slot') or '') == 'central' else 0,
                            wid,
                        ),
                    )
                    conn.commit()
                    msg = 'ذخیره شد'
                else:
                    outside_raw = request.form.get('amount_outside')
                    outside = int(outside_raw) if outside_raw not in (None, '') else None
                    conn.execute(
                        'INSERT INTO wage_rates (work_description, wage_amount, rate_code, amount_outside, is_per_km, category, sort_order, is_active, device_type, install_slot, is_central_device) VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?)',
                        (
                            request.form['work_description'],
                            int(request.form['wage_amount']),
                            request.form.get('rate_code') or None,
                            outside,
                            1 if request.form.get('is_per_km') else 0,
                            request.form.get('category') or None,
                            int(request.form.get('sort_order') or 99),
                            (request.form.get('device_type') or '').strip() or None,
                            (request.form.get('install_slot') or '').strip() or None,
                            1 if (request.form.get('install_slot') or '') == 'central' else 0,
                        ),
                    )
                    conn.commit()
                    msg = 'افزوده شد'
            except Exception as e:
                error = str(e)

        wage_rates = conn.execute(
            'SELECT * FROM wage_rates ORDER BY COALESCE(sort_order, 999), id'
        ).fetchall()
        conn.close()
        try:
            device_types = [r['device_type'] for r in conn.execute(
                'SELECT DISTINCT device_type FROM devices WHERE device_type IS NOT NULL AND device_type != "" ORDER BY device_type'
            ).fetchall()]
        except Exception:
            device_types = []
        return render_template(
            'settings_wages.html', wage_rates=wage_rates,
            active_page='settings', current_user=current_user,
            msg=msg, error=error, device_types=device_types,
        )




    @app.route('/settings/wages/delete/<int:wage_id>', methods=['POST'])
    def delete_wage_rate(wage_id):
        current_user = get_current_user()
        if not current_user or not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')

        conn = get_db()
        conn.execute('DELETE FROM wage_rates WHERE id = ?', (wage_id,))
        conn.commit()
        conn.close()
        return redirect('/settings/wages')





    @app.route('/settings/parts', methods=['GET', 'POST'])
    def settings_parts():
        current_user = get_current_user()
        if not current_user:
            return redirect('/login')
        if not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')

        conn = get_db()
        if request.method == 'POST':
            warehouse_code = request.form.get('warehouse_code', '').strip()
            part_name = request.form['part_name'].strip()
            device_type = request.form.get('device_type', '').strip()
            device_model = request.form.get('device_model', '').strip()
            conn.execute(
                'INSERT INTO parts (warehouse_code, part_name, device_type, device_model) VALUES (?, ?, ?, ?)',
                (warehouse_code or None, part_name, device_type or None, device_model or None)
            )
            conn.commit()
            conn.close()
            return redirect('/settings/parts')

        parts_list = conn.execute('SELECT * FROM parts ORDER BY id DESC').fetchall()
        devices_list = conn.execute('SELECT * FROM devices ORDER BY device_type, model').fetchall()
        conn.close()
        return render_template(
            'settings_parts.html', parts=parts_list, devices=devices_list,
            active_page='settings', current_user=current_user,
            import_error=None, imported_count=None
        )





    @app.route('/settings/parts/delete/<int:part_id>', methods=['POST'])
    def delete_part(part_id):
        current_user = get_current_user()
        if not current_user or not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')
        conn = get_db()
        conn.execute('DELETE FROM warehouse_stock WHERE part_id = ?', (part_id,))
        conn.execute('DELETE FROM parts WHERE id = ?', (part_id,))
        conn.commit()
        conn.close()
        return redirect('/settings/parts')





    @app.route('/settings/parts/import', methods=['POST'])
    def import_parts():
        current_user = get_current_user()
        if not current_user or not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')

        file = request.files.get('excel_file')
        import_error = None
        imported_count = 0

        if not file or file.filename == '':
            import_error = 'هیچ فایلی انتخاب نشده است.'
        elif not file.filename.lower().endswith(('.xlsx', '.xls')):
            import_error = 'فقط فایل اکسل (xlsx یا xls) مجاز است.'
        else:
            try:
                workbook = openpyxl.load_workbook(file, data_only=True)
                sheet = workbook.active
                conn = get_db()
                rows = list(sheet.iter_rows(values_only=True))
                data_rows = rows[1:] if rows else []
                for row in data_rows:
                    if not row or not row[0]:
                        continue
                    values = list(row) + [None] * 4
                    # ترتیب: کد انبار، نام قطعه، نوع دستگاه، مدل دستگاه
                    warehouse_code = values[0]
                    part_name = values[1]
                    device_type = values[2]
                    device_model = values[3]
                    if not part_name:
                        continue
                    conn.execute(
                        'INSERT INTO parts (warehouse_code, part_name, device_type, device_model) VALUES (?, ?, ?, ?)',
                        (str(warehouse_code) if warehouse_code else None,
                         str(part_name),
                         str(device_type) if device_type else None,
                         str(device_model) if device_model else None)
                    )
                    imported_count += 1
                conn.commit()
                conn.close()
            except Exception as e:
                import_error = f'خطا در پردازش فایل: {e}'

        conn = get_db()
        parts_list = conn.execute('SELECT * FROM parts ORDER BY id DESC').fetchall()
        devices_list = conn.execute('SELECT * FROM devices ORDER BY device_type, model').fetchall()
        conn.close()
        return render_template(
            'settings_parts.html', parts=parts_list, devices=devices_list,
            active_page='settings', current_user=current_user,
            import_error=import_error, imported_count=imported_count
        )





    @app.route('/settings/parts/export')
    def export_parts():
        current_user = get_current_user()
        if not current_user or not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')
        conn = get_db()
        parts_list = conn.execute('SELECT * FROM parts ORDER BY id').fetchall()
        conn.close()
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'قطعات'
        sheet.append(['کد انبار', 'نام قطعه', 'نوع دستگاه', 'مدل دستگاه'])
        for p in parts_list:
            sheet.append([p['warehouse_code'], p['part_name'], p['device_type'], p['device_model']])
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='parts.xlsx'
        )




