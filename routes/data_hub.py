# -*- coding: utf-8 -*-
"""هاب یکپارچه‌ی Import / Export."""
from __future__ import annotations
from flask import request, redirect, render_template, url_for, session
import jdatetime
import datetime as _dt
import os
try:
    import openpyxl
except ImportError:
    openpyxl = None

from core.constants import *  # noqa: F401,F403
from core.helpers import *  # noqa: F401,F403


# ---------------------------------------------------------------------------
# فهرست محل‌های Export — هر ماژول به روت اکسپورت موجودش وصل می‌شه
# ---------------------------------------------------------------------------
EXPORT_TARGETS = [
    {'key': 'customers', 'label': 'مشتریان', 'url': '/settings/customers/export', 'perm': 'parties'},
    {'key': 'parts', 'label': 'قطعات', 'url': '/settings/parts/export', 'perm': 'parties'},
    {'key': 'customer_affairs', 'label': 'درخواست‌های امور مشتریان (کلی)', 'url': '/customer-affairs/export', 'perm': 'customer_affairs'},
    {'key': 'internal_repair', 'label': 'تعمیر داخلی', 'url': '/service/internal-repair/export', 'perm': 'customer_affairs'},
    {'key': 'external_repair', 'label': 'تعمیر خارجی', 'url': '/service/external-repair/export', 'perm': 'customer_affairs'},
    {'key': 'installation', 'label': 'نصب و آموزش', 'url': '/service/installation/export', 'perm': 'customer_affairs'},
    {'key': 'inspection', 'label': 'بازدید', 'url': '/service/inspection/export', 'perm': 'customer_affairs'},
    {'key': 'service_performance', 'label': 'گزارش عملکرد خدمات / بایگانی', 'url': '/reports/service-performance?export=1', 'perm': 'reports'},
    {'key': 'management_report', 'label': 'گزارش مدیریتی', 'url': '/reports/management/export', 'perm': 'reports'},
    {'key': 'wage_report', 'label': 'گزارش دستمزد', 'url': '/finance/wage-report/export', 'perm': 'finance'},
]

# فهرست محل‌های Import — فعلاً فقط جاهایی که منطق درج رکورد دارن
IMPORT_TARGETS = [
    {'key': 'customers', 'label': 'مشتریان', 'perm': 'parties',
     'hint': 'ترتیب ستون‌ها: کد طرف حساب، نوع (حقیقی/حقوقی)، نام مشتری، استان، شهرستان، شهر، آدرس، کد پستی، کد ملی/شناسه ملی، شناسه اقتصادی، تلفن ثابت، نام مسئول تجهیزات، تلفن همراه — ستون کد فقط مرجع است. فایل قدیمی بدون ستون نوع هم پذیرفته می‌شود.'},
    {'key': 'parts', 'label': 'قطعات', 'perm': 'parties',
     'hint': 'ترتیب ستون‌ها: کد انبار، نام قطعه، نوع دستگاه، مدل دستگاه'},
    {'key': 'warranty', 'label': 'شناسنامه گارانتی', 'perm': 'parties',
     'hint': 'ترتیب ستون‌ها: نام مشتری، نوع دستگاه، مدل، سریال، تاریخ تولید، شهریار (سریال تکراری رد می‌شود)'},
    {'key': 'internal_repair', 'label': 'تعمیر داخلی (بایگانی)', 'perm': 'customer_affairs',
     'hint': 'ترتیب ستون‌ها: نام مشتری، آدرس، تلفن، نوع دستگاه، سریال، تاریخ نصب، شرح مشکل، استان، تکنسین مسئول، تاریخ پذیرش'},
    {'key': 'external_repair', 'label': 'تعمیر خارجی (بایگانی)', 'perm': 'customer_affairs',
     'hint': 'ترتیب ستون‌ها: نام مشتری، تلفن، نوع دستگاه، سریال، استان، شهر، تکنسین مسئول، شرح درخواست، تاریخ پذیرش، تاریخ پایان گارانتی'},
    {'key': 'installation', 'label': 'نصب و آموزش (بایگانی)', 'perm': 'customer_affairs',
     'hint': 'ترتیب ستون‌ها: نام مشتری، تلفن، آدرس، استان، شهر، تکنسین مسئول، یادداشت نصب، تاریخ پذیرش'},
    {'key': 'inspection', 'label': 'بازدید (بایگانی)', 'perm': 'customer_affairs',
     'hint': 'ترتیب ستون‌ها: نام مشتری، تلفن، استان، شهر، نوع دستگاه بازدید، تعداد دستگاه، تکنسین مسئول، شرح درخواست، تاریخ پذیرش، بازدید پولی(بله/خیر)'},
    {'key': 'geo', 'label': 'لیست استان / شهرستان / شهر', 'perm': 'system',
     'hint': 'ستون‌ها: استان، شهرستان، شهر — فایل Iran_Provinces_Counties_Cities.xlsx قابل استفاده است'},
]


INCREMENTAL_EXCEL_TARGET = {
    'key': 'incremental_excel',
    'label': 'Excel روزانه (مشتری + سرویس)',
    'perm': 'customer_affairs',
    'hint': 'برای خروجی‌های روزانه نرم‌افزار مبدأ. کد طرف حساب = شناسه مشتری و شماره پذیرش = شناسه سرویس؛ رکوردهای جدید/تغییرکرده/تکراری جدا می‌شوند.'
}


def _allowed_targets(targets, current_user):
    out = []
    for t in targets:
        if _is_reception(current_user) or user_has_access(current_user, t['perm']):
            out.append(t)
    return out


def register(app):

    @app.route('/settings/data-hub')
    def data_hub():
        current_user = get_current_user()
        if not current_user:
            return redirect('/login')

        export_targets = _allowed_targets(EXPORT_TARGETS, current_user)
        import_targets = _allowed_targets(IMPORT_TARGETS, current_user)

        return render_template(
            'settings_data_hub.html',
            export_targets=export_targets,
            import_targets=import_targets,
            active_page='data-hub',
            current_user=current_user,
            result=request.args.get('result'),
            result_module=request.args.get('module'),
            result_count=request.args.get('count'),
            result_msg=request.args.get('msg'),
            incremental_preview=session.get('excel_incremental_preview'),
            incremental_result=session.pop('excel_incremental_result', None),
        )

    @app.route('/settings/data-hub/export')
    def data_hub_export():
        current_user = get_current_user()
        if not current_user:
            return redirect('/login')
        module = request.args.get('module', '')
        fmt = request.args.get('format', 'excel')
        target = next((t for t in EXPORT_TARGETS if t['key'] == module), None)
        if not target or not (_is_reception(current_user) or user_has_access(current_user, target['perm'])):
            return redirect('/settings/data-hub?result=error&msg=' + 'دسترسی یا محل نامعتبر')

        if fmt == 'pdf':
            # خروجی PDF هنوز پیاده‌سازی نشده (نیازمند فونت فارسی/راست‌به‌چپ)
            return redirect('/settings/data-hub?result=error&module=' + module +
                             '&msg=' + 'خروجی PDF هنوز آماده نیست — فعلاً فقط اکسل')

        sep = '&' if '?' in target['url'] else '?'
        return redirect(target['url'] + sep + '_from_hub=1')

    @app.route('/settings/data-hub/incremental-excel/preview', methods=['POST'])
    def incremental_excel_preview():
        current_user = get_current_user()
        if not current_user:
            return redirect('/login')
        if not (_is_reception(current_user) or user_has_access(current_user, 'customer_affairs')):
            return redirect('/settings/data-hub?result=error&msg=دسترسی کافی نیست')

        file = request.files.get('incremental_excel_file')
        if not file or not file.filename:
            return redirect('/settings/data-hub?result=error&msg=فایل Excel انتخاب نشده است')
        if not file.filename.lower().endswith('.xlsx'):
            return redirect('/settings/data-hub?result=error&msg=فقط فایل xlsx مجاز است')

        try:
            from core.excel_incremental import save_upload, analyze
            path = save_upload(file)
            report = analyze(path)
            token = os.path.basename(path)
            # Only retain the server-side temporary path, never the workbook itself in session.
            session['excel_incremental_preview'] = {
                'path': path,
                'token': token,
                'file': report['file'],
                'customers': report['customers'],
                'services': report['services'],
                'total_customers': report['total_customers'],
                'total_services': report['total_services'],
            }
            return redirect('/settings/data-hub#incremental-excel')
        except Exception as e:
            return redirect('/settings/data-hub?result=error&msg=' + str(e))

    @app.route('/settings/data-hub/incremental-excel/commit', methods=['POST'])
    def incremental_excel_commit():
        current_user = get_current_user()
        if not current_user:
            return redirect('/login')
        if not (_is_reception(current_user) or user_has_access(current_user, 'customer_affairs')):
            return redirect('/settings/data-hub?result=error&msg=دسترسی کافی نیست')

        preview = session.get('excel_incremental_preview') or {}
        path = preview.get('path')
        if not path or not os.path.isfile(path):
            return redirect('/settings/data-hub?result=error&msg=پیش‌نمایش منقضی شده است؛ فایل را دوباره انتخاب کنید')

        try:
            from core.excel_incremental import commit, cleanup_upload
            result = commit(path)
            cleanup_upload(path)
            session.pop('excel_incremental_preview', None)
            session['excel_incremental_result'] = {
                'status': result.get('status'),
                'file': result.get('file'),
                'customers': result.get('customers'),
                'services': result.get('services'),
                'batch_id': result.get('batch_id'),
                'backup': result.get('backup'),
            }
            return redirect('/settings/data-hub#incremental-excel')
        except Exception as e:
            return redirect('/settings/data-hub?result=error&msg=' + str(e))

    @app.route('/settings/data-hub/incremental-excel/cancel', methods=['POST'])
    def incremental_excel_cancel():
        current_user = get_current_user()
        if not current_user:
            return redirect('/login')
        preview = session.pop('excel_incremental_preview', None) or {}
        path = preview.get('path')
        if path:
            try:
                from core.excel_incremental import cleanup_upload
                cleanup_upload(path)
            except Exception:
                pass
        return redirect('/settings/data-hub#incremental-excel')

    @app.route('/settings/data-hub/import', methods=['POST'])
    def data_hub_import():
        current_user = get_current_user()
        if not current_user:
            return redirect('/login')
        module = request.form.get('module', '')
        target = next((t for t in IMPORT_TARGETS if t['key'] == module), None)
        if not target or not (_is_reception(current_user) or user_has_access(current_user, target['perm'])):
            return redirect('/settings/data-hub?result=error&msg=' + 'دسترسی یا محل نامعتبر')

        file = request.files.get('excel_file')
        if not file or file.filename == '':
            return redirect('/settings/data-hub?result=error&module=' + module + '&msg=' + 'هیچ فایلی انتخاب نشده است')
        # فقط xlsx — openpyxl از xls پشتیبانی نمی‌کند (xlsx = zip)
        if not file.filename.lower().endswith('.xlsx'):
            return redirect('/settings/data-hub?result=error&module=' + module + '&msg=' + 'فقط فایل اکسل با فرمت xlsx مجاز است')

        try:
            handlers = {
                'customers': _import_customers_rows,
                'parts': _import_parts_rows,
                'warranty': _import_warranty_rows,
                'internal_repair': _import_internal_repair_rows,
                'external_repair': _import_external_repair_rows,
                'installation': _import_installation_rows,
                'inspection': _import_inspection_rows,
                'geo': _import_geo_rows,
            }
            if module not in handlers:
                return redirect('/settings/data-hub?result=error&module=' + module + '&msg=' + 'محل نامعتبر')
            count = handlers[module](file)
        except Exception as e:
            msg = str(e)
            if 'not a zip file' in msg.lower() or type(e).__name__ == 'BadZipFile':
                msg = 'فایل اکسل معتبر نیست. لطفاً فقط فایل با فرمت xlsx آپلود کنید.'
            return redirect('/settings/data-hub?result=error&module=' + module + '&msg=' + f'خطا در پردازش فایل: {msg}')

        return redirect(f'/settings/data-hub?result=ok&module={module}&count={count}')


# ---------------------------------------------------------------------------
# منطق درج ردیف‌ها — همون منطق فرم‌های اختصاصی، برای استفاده‌ی مشترک از هاب
# ---------------------------------------------------------------------------
def _import_customers_rows(file):
    """ایمپورت مشتریان با تشخیص ستون از روی هدر (ترجیحی) یا ترتیب ثابت.
    هدرهای شناخته‌شده: کد، نوع، نام، استان، شهرستان، شهر، آدرس، کد پستی،
    کد ملی، شناسه اقتصادی، تلفن، مسئول تجهیزات، موبایل.
    """
    from core.geo import ensure_geo_schema

    def _cell(val):
        if val is None:
            return None
        s = str(val).strip()
        if not s or s.lower() in ('none', 'null', '-'):
            return None
        return s

    def _norm_type(val):
        s = _cell(val)
        if not s:
            return None
        s2 = s.replace('ي', 'ی').replace('ك', 'ک')
        if s2 in ('حقیقی', 'حقوقی'):
            return s2
        low = s2.lower()
        if low in ('haghighi', 'natural', 'individual'):
            return 'حقیقی'
        if low in ('hoghooghi', 'legal', 'company'):
            return 'حقوقی'
        return None

    def _norm_header(h):
        s = str(h or '').strip().replace('ي', 'ی').replace('ك', 'ک').replace('‌', '')
        s = s.replace(' ', '').replace('_', '').lower()
        return s

    # نگاشت هدر → فیلد
    HEADER_MAP = {
        'کدطرفحساب': '_code',
        'کد': '_code',
        'id': '_code',
        'نوع': 'person_type',
        'نوعطرفحساب': 'person_type',
        'persontype': 'person_type',
        'ناممشتری': 'name',
        'نام': 'name',
        'نامشرکت': 'name',
        'مشتری': 'name',
        'استان': 'province',
        'شهرستان': 'county',
        'شهر': 'city',
        'آدرس': 'address',
        'نشانی': 'address',
        'کدپستی': 'postal_code',
        'کدملی': 'national_id',
        'شناسهملی': 'national_id',
        'کدملی/شناسهملی': 'national_id',
        'کد/شناسهملی': 'national_id',
        'شناسهاقتصادی': 'economic_code',
        'کداقتصادی': 'economic_code',
        'تلفنثابت': 'phone',
        'تلفن': 'phone',
        'ناممسئولتجهیزات': 'equipment_manager_name',
        'مسئولتجهیزات': 'equipment_manager_name',
        'تلفنهمراه': 'equipment_manager_mobile',
        'موبایل': 'equipment_manager_mobile',
        'موبایلمسئول': 'equipment_manager_mobile',
        'شمارهموبایلمسئولتجهیزات': 'equipment_manager_mobile',
    }

    workbook = openpyxl.load_workbook(file, data_only=True)
    sheet = workbook.active
    conn = get_db()
    ensure_geo_schema(conn)
    cols = {row[1] for row in conn.execute('PRAGMA table_info(customers)').fetchall()}
    for col, default in (
        ('person_type', "TEXT DEFAULT 'حقیقی'"),
        ('county', 'TEXT'),
        ('postal_code', 'TEXT'),
    ):
        if col not in cols:
            try:
                conn.execute('ALTER TABLE customers ADD COLUMN ' + col + ' ' + default)
            except Exception:
                pass

    all_rows = list(sheet.iter_rows(values_only=True))
    if not all_rows:
        conn.close()
        return 0

    header_row = all_rows[0]
    field_index = {}
    for i, h in enumerate(header_row):
        key = HEADER_MAP.get(_norm_header(h))
        if key and key not in field_index:
            field_index[key] = i

    use_headers = 'name' in field_index
    data_rows = all_rows[1:]
    count = 0

    for row in data_rows:
        if not row:
            continue
        values = list(row) + [None] * 20

        if use_headers:
            def g(field):
                idx = field_index.get(field)
                if idx is None or idx >= len(values):
                    return None
                return values[idx]

            name = _cell(g('name'))
            person_type = _norm_type(g('person_type')) or 'حقیقی'
            province = _cell(g('province'))
            county = _cell(g('county'))
            city = _cell(g('city'))
            address = _cell(g('address'))
            postal_code = _cell(g('postal_code'))
            national_id = _cell(g('national_id'))
            economic_code = _cell(g('economic_code'))
            phone = _cell(g('phone'))
            eqm_name = _cell(g('equipment_manager_name'))
            eqm_mobile = _cell(g('equipment_manager_mobile'))
        else:
            # ترتیب ثابت: کد | نوع | نام | استان | شهرستان | شهر | آدرس | کدپستی | ...
            v1 = _cell(values[1])
            ptype = _norm_type(v1)
            if ptype:
                person_type = ptype
                name = _cell(values[2])
                province = _cell(values[3])
                county = _cell(values[4])
                city = _cell(values[5])
                address = _cell(values[6])
                postal_code = _cell(values[7])
                national_id = _cell(values[8])
                economic_code = _cell(values[9])
                phone = _cell(values[10])
                eqm_name = _cell(values[11])
                eqm_mobile = _cell(values[12])
            else:
                if values[1] is None:
                    continue
                person_type = 'حقیقی'
                name = _cell(values[1])
                province = _cell(values[2])
                county = _cell(values[3])
                city = _cell(values[4])
                address = _cell(values[5])
                postal_code = _cell(values[6])
                national_id = _cell(values[7])
                economic_code = _cell(values[8])
                phone = _cell(values[9])
                eqm_name = _cell(values[10])
                eqm_mobile = _cell(values[11])

        if not name or name in ('نام مشتری', 'نام'):
            continue

        conn.execute(
            "INSERT INTO customers "
            "(name, phone, national_id, economic_code, province, county, city, address, postal_code, "
            "equipment_manager_name, equipment_manager_mobile, person_type) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                name, phone, national_id, economic_code, province, county, city, address,
                postal_code, eqm_name, eqm_mobile, person_type,
            ),
        )
        count += 1

    conn.commit()
    conn.close()
    return count


def _import_parts_rows(file):
    workbook = openpyxl.load_workbook(file, data_only=True)
    sheet = workbook.active
    conn = get_db()
    rows = list(sheet.iter_rows(values_only=True))
    data_rows = rows[1:] if rows else []
    count = 0
    for row in data_rows:
        if not row or not row[0]:
            continue
        values = list(row) + [None] * 4
        # ترتیب: کد انبار، نام قطعه، نوع دستگاه، مدل دستگاه
        warehouse_code, part_name, device_type, device_model = values[:4]
        if not part_name:
            continue
        conn.execute(
            'INSERT INTO parts (warehouse_code, part_name, device_type, device_model) VALUES (?, ?, ?, ?)',
            (str(warehouse_code) if warehouse_code else None, str(part_name),
             str(device_type) if device_type else None, str(device_model) if device_model else None)
        )
        count += 1
    conn.commit()
    conn.close()
    return count


def _jtoday():
    return jdatetime.date.today().strftime('%Y/%m/%d')


def _norm_date(value):
    """نرمال‌سازی مقدار تاریخ از سلول اکسل به رشته‌ی شمسی YYYY/MM/DD.
    اکسل گاهی تاریخ رو به‌صورت واقعی (datetime/date میلادی) می‌ده، نه متن —
    اگر مستقیم str() بشه، یه رشته‌ی میلادی مثل '2025-05-15 00:00:00' ذخیره
    می‌شه که با فرمت شمسی مورد انتظار بقیه‌ی اپ (مثل تاریخ پذیرش/نصب) جور
    درنمی‌آد و فیلتر/مرتب‌سازی تاریخ رو خراب می‌کنه. این تابع اون حالت رو
    تشخیص می‌ده و به شمسی تبدیل می‌کنه؛ رشته‌های متنی (از قبل شمسی) دست‌نخورده می‌مونن.
    """
    if value in (None, ''):
        return None
    if isinstance(value, _dt.datetime):
        value = value.date()
    if isinstance(value, _dt.date):
        try:
            return jdatetime.date.fromgregorian(date=value).strftime('%Y/%m/%d')
        except Exception:
            return str(value)
    return str(value)


def _import_internal_repair_rows(file):
    workbook = openpyxl.load_workbook(file, data_only=True)
    sheet = workbook.active
    conn = get_db()
    from core.reception_numbers import assign_reception_no
    rows = list(sheet.iter_rows(values_only=True))
    data_rows = rows[1:] if rows else []
    count = 0
    cfg = SERVICE_CATEGORIES['internal-repair']
    for row in data_rows:
        if not row or not row[0]:
            continue
        values = list(row) + [None] * 10
        (customer_name, customer_address, customer_phone, device_type, serial_number,
         installation_date, problem_desc, province, assigned_technician, reception_date) = values[:10]
        reception_date = _norm_date(reception_date) or _jtoday()
        installation_date = _norm_date(installation_date)
        cur = conn.execute(
            '''INSERT INTO requests
               (customer_name, customer_address, customer_phone, device_type, serial_number,
                installation_date, problem_desc, service_type, province, assigned_technician,
                reception_date, request_category)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
            (str(customer_name), str(customer_address) if customer_address else '',
             str(customer_phone) if customer_phone else '', str(device_type) if device_type else '',
             str(serial_number) if serial_number else '', installation_date or '',
             str(problem_desc) if problem_desc else '', cfg['service_type'],
             str(province) if province else '', str(assigned_technician) if assigned_technician else None,
             reception_date, cfg['category'])
        )
        new_id = cur.lastrowid
        try:
            assign_reception_no(conn, new_id, request_category=cfg['category'])
        except Exception:
            pass
        count += 1
    conn.commit()
    conn.close()
    return count


def _import_external_repair_rows(file):
    workbook = openpyxl.load_workbook(file, data_only=True)
    sheet = workbook.active
    conn = get_db()
    from core.reception_numbers import assign_reception_no
    rows = list(sheet.iter_rows(values_only=True))
    data_rows = rows[1:] if rows else []
    count = 0
    for row in data_rows:
        if not row or not row[0]:
            continue
        values = list(row) + [None] * 10
        (customer_name, customer_phone, device_type, serial_number, province, city,
         assigned_technician, hospital_request_desc, reception_date, warranty_end_date) = values[:10]
        reception_date = _norm_date(reception_date) or _jtoday()
        warranty_end_date = _norm_date(warranty_end_date)
        problem_desc = str(hospital_request_desc) if hospital_request_desc else 'تعمیر خارجی'
        cur = conn.execute(
            '''INSERT INTO requests
               (customer_name, device_type, serial_number, problem_desc, service_type, province,
                assigned_technician, reception_date, request_category, status, city,
                hospital_request_desc, warranty_end_date, customer_phone)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
            (str(customer_name), str(device_type) if device_type else '—',
             str(serial_number) if serial_number else '', problem_desc, 'در محل',
             str(province) if province else '', str(assigned_technician) if assigned_technician else None,
             reception_date, 'تعمیر', EXTERNAL_REPAIR_STATUSES[1], str(city) if city else '',
             str(hospital_request_desc) if hospital_request_desc else '',
             warranty_end_date or '', str(customer_phone) if customer_phone else '')
        )
        new_id = cur.lastrowid
        try:
            assign_reception_no(conn, new_id)
        except Exception:
            pass
        if assigned_technician:
            try:
                sync_request_technician(
                    conn, new_id, assigned_technician,
                    expected_role='تکنسین استانی', province=str(province or '').strip() or None,
                )
            except Exception as exc:
                conn.rollback()
                conn.close()
                raise ValueError(f'تکنسین استانی نامعتبر برای ردیف {customer_name}: {exc}') from exc
        if reception_date:
            try:
                _set_stage_date(conn, new_id, EXTERNAL_REPAIR_STATUSES[1], reception_date, 'ورود از فایل بایگانی')
            except Exception:
                pass
        count += 1
    conn.commit()
    conn.close()
    return count


def _import_installation_rows(file):
    workbook = openpyxl.load_workbook(file, data_only=True)
    sheet = workbook.active
    conn = get_db()
    from core.reception_numbers import assign_reception_no
    rows = list(sheet.iter_rows(values_only=True))
    data_rows = rows[1:] if rows else []
    count = 0
    for row in data_rows:
        if not row or not row[0]:
            continue
        values = list(row) + [None] * 8
        (customer_name, customer_phone, customer_address, province, city,
         assigned_technician, install_notes, reception_date) = values[:8]
        reception_date = _norm_date(reception_date) or _jtoday()
        problem_desc = str(install_notes) if install_notes else 'درخواست نصب و آموزش'
        cur = conn.execute(
            '''INSERT INTO requests
               (customer_name, customer_address, customer_phone, device_type, serial_number,
                problem_desc, service_type, province, assigned_technician,
                reception_date, request_category, status, city, install_notes, request_subtype)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
            (str(customer_name), str(customer_address) if customer_address else '',
             str(customer_phone) if customer_phone else '', 'نصب و آموزش', None,
             problem_desc, 'در محل', str(province) if province else '',
             str(assigned_technician) if assigned_technician else None,
             reception_date, 'نصب و آموزش', INSTALLATION_REQUEST_STATUSES[0],
             str(city) if city else '', str(install_notes) if install_notes else '', 'درخواست نصب')
        )
        new_id = cur.lastrowid
        try:
            assign_reception_no(conn, new_id)
        except Exception:
            pass
        if reception_date:
            try:
                _set_stage_date(conn, new_id, INSTALLATION_REQUEST_STATUSES[0], reception_date, 'ورود از فایل بایگانی')
            except Exception:
                pass
        count += 1
    conn.commit()
    conn.close()
    return count


def _import_inspection_rows(file):
    workbook = openpyxl.load_workbook(file, data_only=True)
    sheet = workbook.active
    conn = get_db()
    from core.reception_numbers import assign_reception_no
    rows = list(sheet.iter_rows(values_only=True))
    data_rows = rows[1:] if rows else []
    count = 0
    for row in data_rows:
        if not row or not row[0]:
            continue
        values = list(row) + [None] * 10
        (customer_name, customer_phone, province, city, visit_device_type, visit_device_count,
         assigned_technician, hospital_request_desc, reception_date, paid_flag) = values[:10]
        reception_date = _norm_date(reception_date) or _jtoday()
        problem_desc = str(hospital_request_desc) if hospital_request_desc else 'درخواست بازدید'
        is_paid_visit = 1 if str(paid_flag or '').strip() in ('بله', '1', 'yes', 'Yes') else 0
        initial_status = 'در انتظار ارسال پیش‌فاکتور است' if is_paid_visit else 'در انتظار انجام بازدید'
        cur = conn.execute(
            '''INSERT INTO requests
               (customer_name, customer_phone, device_type, problem_desc, service_type, province,
                assigned_technician, reception_date, request_category, status, city,
                hospital_request_desc, is_paid_visit, visit_device_type, visit_device_count, request_subtype)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
            (str(customer_name), str(customer_phone) if customer_phone else '',
             str(visit_device_type) if visit_device_type else 'بازدید', problem_desc, 'در محل',
             str(province) if province else '', str(assigned_technician) if assigned_technician else None,
             reception_date, 'بررسی و بازدید', initial_status, str(city) if city else '',
             str(hospital_request_desc) if hospital_request_desc else '', is_paid_visit,
             str(visit_device_type) if visit_device_type else '',
             int(visit_device_count) if visit_device_count not in (None, '') else None, 'درخواست بازدید')
        )
        new_id = cur.lastrowid
        try:
            assign_reception_no(conn, new_id)
        except Exception:
            pass
        if reception_date:
            try:
                _set_stage_date(conn, new_id, 'درخواست ثبت شد', reception_date, 'ورود از فایل بایگانی')
            except Exception:
                pass
        count += 1
    conn.commit()
    conn.close()
    return count


def _import_warranty_rows(file):
    workbook = openpyxl.load_workbook(file, data_only=True)
    sheet = workbook.active
    conn = get_db()
    rows = list(sheet.iter_rows(values_only=True))
    data_rows = rows[1:] if rows else []
    count = 0
    for row in data_rows:
        if not row:
            continue
        values = list(row) + [None] * 8
        customer_name, device_type, model, serial_number, production_date = values[:5]
        shahriar = values[5] if len(values) > 5 else None
        if not serial_number:
            continue
        existing = conn.execute(
            'SELECT id FROM warranty_cards WHERE serial_number = ?', (str(serial_number),)
        ).fetchone()
        if existing:
            continue
        conn.execute(
            "INSERT INTO warranty_cards "
            "(customer_name, device_type, model, serial_number, production_date, "
            "installation_date, shahriar, warranty_status) "
            "VALUES (?, ?, ?, ?, ?, NULL, ?, 'تولید شده')",
            (str(customer_name) if customer_name else None, str(device_type) if device_type else None,
             str(model) if model else None, str(serial_number), _norm_date(production_date),
             str(shahriar) if shahriar else None)
        )
        count += 1
    conn.commit()
    conn.close()
    return count


def _import_geo_rows(file):
    """ورود لیست استان/شهرستان/شهر از اکسل — همان منطق settings/geo."""
    from core.geo import import_geo_from_excel, ensure_geo_schema
    conn = get_db()
    ensure_geo_schema(conn)
    # clear_existing پیش‌فرض True (مثل فرم قبلی)
    p, c, ci, err = import_geo_from_excel(conn, file, clear_existing=True)
    if err:
        conn.close()
        raise RuntimeError(err)
    conn.commit()
    conn.close()
    return (p or 0) + (c or 0) + (ci or 0)

