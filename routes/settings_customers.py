# -*- coding: utf-8 -*-
"""مسیرهای طرف‌حساب‌ها (customers): لیست/ستون‌ها/ویرایش/مشاهده/حذف/خروجی/ورودی."""
# -*- coding: utf-8 -*-
from __future__ import annotations

import logging

from flask import (
    request, session, redirect, render_template, jsonify, send_file, current_app
)

logger = logging.getLogger("sazgan")
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from io import BytesIO
import os, re, sqlite3, base64, time
import jdatetime
try:
    import openpyxl
except ImportError:
    openpyxl = None

from core.constants import *
from core import helpers as H


from core.helpers import *  # noqa: F401,F403
from core.constants import *  # noqa: F401,F403



# ---------------------------------------------------------------------------
# ستون‌های لیست طرف حساب‌ها — تنظیم سراسری فقط توسط ادمین
# ---------------------------------------------------------------------------
CUSTOMER_LIST_COLUMN_DEFS = [
    {'key': 'actions', 'label': 'عملیات', 'required': True},
    {'key': 'id', 'label': 'کد', 'required': True},
    {'key': 'person_type', 'label': 'نوع', 'required': False},
    {'key': 'name', 'label': 'نام مشتری', 'required': True},
    {'key': 'province', 'label': 'استان', 'required': False},
    {'key': 'county', 'label': 'شهرستان', 'required': False},
    {'key': 'city', 'label': 'شهر', 'required': False},
    {'key': 'address', 'label': 'آدرس', 'required': False},
    {'key': 'postal_code', 'label': 'کد پستی', 'required': False},
    {'key': 'national_id', 'label': 'کد/شناسه ملی', 'required': False},
    {'key': 'economic_code', 'label': 'شناسه اقتصادی', 'required': False},
    {'key': 'phone', 'label': 'تلفن ثابت', 'required': False},
    {'key': 'equipment_manager_name', 'label': 'مسئول تجهیزات', 'required': False},
    {'key': 'equipment_manager_mobile', 'label': 'موبایل مسئول', 'required': False},
]
CUSTOMER_LIST_COLUMNS_KEY = 'customers_list_columns'
_CUSTOMER_COL_KEYS = {c['key'] for c in CUSTOMER_LIST_COLUMN_DEFS}
_CUSTOMER_COL_REQUIRED = {c['key'] for c in CUSTOMER_LIST_COLUMN_DEFS if c.get('required')}


def _default_customer_columns():
    return [c['key'] for c in CUSTOMER_LIST_COLUMN_DEFS]


def _load_customer_columns(conn=None):
    """لیست کلید ستون‌های فعال به ترتیب. همیشه actions/id/name را دارد."""
    import json as _json
    raw = None
    own_conn = False
    try:
        if conn is None:
            conn = get_db()
            own_conn = True
        raw = get_setting(conn, CUSTOMER_LIST_COLUMNS_KEY, '') or ''
    except Exception:
        raw = ''
    finally:
        if own_conn and conn is not None:
            try:
                conn.close()
            except Exception as e:
                logger.exception("Silent exception in routes/settings.py")
    keys = None
    if raw:
        try:
            data = _json.loads(raw)
            if isinstance(data, dict):
                keys = data.get('visible') or data.get('order') or data.get('columns')
            elif isinstance(data, list):
                keys = data
        except Exception:
            keys = None
    if not keys or not isinstance(keys, list):
        keys = _default_customer_columns()
    # فقط کلیدهای معتبر + اجباری‌ها
    out = []
    seen = set()
    for k in keys:
        if k in _CUSTOMER_COL_KEYS and k not in seen:
            out.append(k)
            seen.add(k)
    for k in _default_customer_columns():
        if k in _CUSTOMER_COL_REQUIRED and k not in seen:
            # اجباری‌ها را در موقعیت پیش‌فرض درج کن
            if k == 'actions':
                out.insert(0, k)
            elif k == 'id':
                # بعد از actions
                idx = out.index('actions') + 1 if 'actions' in out else 0
                out.insert(idx, k)
            elif k == 'name':
                # بعد از person_type یا id
                if 'person_type' in out:
                    out.insert(out.index('person_type') + 1, k)
                elif 'id' in out:
                    out.insert(out.index('id') + 1, k)
                else:
                    out.append(k)
            else:
                out.append(k)
            seen.add(k)
    return out


def _customer_columns_for_template(conn=None):
    """لیست دیکشنری ستون برای قالب: key, label, required, visible."""
    visible = set(_load_customer_columns(conn))
    cols = []
    for c in CUSTOMER_LIST_COLUMN_DEFS:
        cols.append({
            'key': c['key'],
            'label': c['label'],
            'required': bool(c.get('required')),
            'visible': c['key'] in visible or bool(c.get('required')),
        })
    # ترتیب نمایش = ترتیب visible keys
    order = _load_customer_columns(conn)
    ordered = []
    by_key = {c['key']: c for c in cols}
    for k in order:
        if k in by_key and by_key[k]['visible']:
            ordered.append(by_key[k])
    return ordered


def _can_edit_customer_columns(user):
    if not user:
        return False
    role = (user.get('role') if isinstance(user, dict) else None) or ''
    return role in ('مدیر سیستم', 'مدیر')




def register(app):
    @app.route('/settings/customers', methods=['GET', 'POST'])
    def settings_customers():
        """طرف حساب‌ها — مهاجرت خودکار + نمایش خطای واضح."""
        try:
            current_user = get_current_user()
            if not current_user:
                return redirect('/login')
            role = current_user['role']
            if role not in ('مسئول پذیرش', 'مدیر سیستم', 'مدیر'):
                return redirect('/')

            conn = get_db()
            conn.execute(
                "CREATE TABLE IF NOT EXISTS customers ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "name TEXT NOT NULL,"
                "phone TEXT,"
                "national_id TEXT,"
                "economic_code TEXT,"
                "province TEXT,"
                "city TEXT,"
                "address TEXT,"
                "equipment_manager_name TEXT,"
                "equipment_manager_mobile TEXT)"
            )
            cust_cols = {row[1] for row in conn.execute('PRAGMA table_info(customers)').fetchall()}
            for col in (
                'phone', 'national_id', 'economic_code', 'province', 'county', 'city', 'address',
                'postal_code', 'equipment_manager_name', 'equipment_manager_mobile',
            ):
                if col not in cust_cols:
                    try:
                        conn.execute('ALTER TABLE customers ADD COLUMN ' + col + ' TEXT')
                    except Exception as e:
                        logger.exception("Silent exception in routes/settings.py")
            if 'person_type' not in cust_cols:
                try:
                    conn.execute("ALTER TABLE customers ADD COLUMN person_type TEXT DEFAULT 'حقیقی'")
                except Exception as e:
                    logger.exception("Silent exception in routes/settings.py")
            conn.commit()

            def _rows():
                out = []
                for r in conn.execute('SELECT * FROM customers WHERE IFNULL(is_archived,0)=0 ORDER BY id DESC').fetchall():
                    d = dict(r)
                    for k in (
                        'id', 'name', 'phone', 'national_id', 'economic_code',
                        'province', 'county', 'city', 'address', 'postal_code',
                        'equipment_manager_name', 'equipment_manager_mobile', 'person_type',
                    ):
                        d.setdefault(k, None)
                    out.append(d)
                return out

            if request.method == 'POST':
                name = (request.form.get('name') or '').strip()
                city_val = (request.form.get('city') or request.form.get('city_manual') or '').strip()
                provinces_list = (get_geo_provinces(conn) if get_geo_provinces else PROVINCES)
                if not name or not city_val:
                    from core.pagination import parse_page, parse_per_page, paginate, pagination_context
                    all_rows = _rows()
                    page = parse_page(request.args)
                    per_page = parse_per_page(request.args, 25)
                    page_result = paginate(all_rows, page=page, per_page=per_page)
                    pctx = pagination_context(page_result, request.args)
                    visible_cols = _customer_columns_for_template(conn)
                    can_cols = _can_edit_customer_columns(current_user)
                    conn.close()
                    return render_template(
                        'settings_customers.html',
                        customers=page_result['items'],
                        provinces=provinces_list,
                        active_page='settings',
                        current_user=current_user,
                        import_error='نام طرف‌حساب و شهر الزامی است.',
                        imported_count=None,
                        customer_columns=visible_cols,
                        can_edit_columns=can_cols,
                        all_column_defs=CUSTOMER_LIST_COLUMN_DEFS,
                        **pctx,
                    )
                person_type = (request.form.get('person_type') or 'حقیقی').strip()
                if person_type not in ('حقیقی', 'حقوقی'):
                    person_type = 'حقیقی'
                conn.execute(
                    'INSERT INTO customers '
                    '(name, phone, national_id, economic_code, province, county, city, address, postal_code, '
                    'equipment_manager_name, equipment_manager_mobile, person_type) '
                    'VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)',
                    (
                        name,
                        request.form.get('phone'),
                        request.form.get('national_id'),
                        request.form.get('economic_code'),
                        request.form.get('province'),
                        request.form.get('county') or None,
                        request.form.get('city') or request.form.get('city_manual'),
                        request.form.get('address'),
                        request.form.get('postal_code'),
                        request.form.get('equipment_manager_name'),
                        request.form.get('equipment_manager_mobile'),
                        person_type,
                    ),
                )
                conn.commit()
                conn.close()
                return redirect('/settings/customers')

            # GET — صفحه‌بندی ۲۵تایی + جستجو سمت سرور
            from core.pagination import parse_page, parse_per_page, paginate, pagination_context
            provinces_list = (get_geo_provinces(conn) if get_geo_provinces else PROVINCES)
            all_rows = _rows()
            q = (request.args.get('q') or '').strip()
            if q:
                ql = q.lower()
                def _match(c):
                    parts = [
                        str(c.get('id') or ''),
                        str(c.get('name') or ''),
                        str(c.get('phone') or ''),
                        str(c.get('province') or ''),
                        str(c.get('county') or ''),
                        str(c.get('city') or ''),
                        str(c.get('national_id') or ''),
                        str(c.get('equipment_manager_name') or ''),
                        str(c.get('equipment_manager_mobile') or ''),
                    ]
                    return any(ql in p.lower() for p in parts if p)
                all_rows = [c for c in all_rows if _match(c)]
            page = parse_page(request.args)
            per_page = parse_per_page(request.args, 25)
            page_result = paginate(all_rows, page=page, per_page=per_page)
            pctx = pagination_context(page_result, request.args)
            visible_cols = _customer_columns_for_template(conn)
            can_cols = _can_edit_customer_columns(current_user)
            conn.close()
            return render_template(
                'settings_customers.html',
                customers=page_result['items'],
                provinces=provinces_list,
                active_page='settings',
                current_user=current_user,
                import_error=None,
                imported_count=None,
                search_q=q,
                customer_columns=visible_cols,
                can_edit_columns=can_cols,
                all_column_defs=CUSTOMER_LIST_COLUMN_DEFS,
                **pctx,
            )
        except Exception as e:
            import traceback
            tb = traceback.format_exc()
            body = (
                '<!DOCTYPE html><html lang="fa" dir="rtl"><meta charset="utf-8">'
                '<title>خطای طرف حساب‌ها</title>'
                '<body style="font-family:Tahoma;padding:24px;background:#fff1f2;color:#9f1239">'
                '<h2>خطا در صفحه طرف حساب‌ها</h2>'
                '<p>متن زیر را کپی کرده و ارسال کنید:</p>'
                '<pre style="background:#fff;border:1px solid #fecdd3;padding:16px;overflow:auto;'
                'direction:ltr;text-align:left;font-size:12px;color:#111">'
                + (type(e).__name__ + ': ' + str(e) + '\n\n' + tb).replace('<', '&lt;')
                + '</pre><p><a href="/">بازگشت</a></p></body></html>'
            )
            return body, 500






    @app.route('/settings/customers/columns', methods=['POST'])
    def settings_customers_columns():
        """ذخیره ستون‌های لیست طرف حساب‌ها — فقط مدیر / مدیر سیستم."""
        import json as _json
        current_user = get_current_user()
        if not current_user or not _can_edit_customer_columns(current_user):
            return redirect('/settings/customers')
        if not _validate_csrf():
            return 'CSRF', 400
        # ترتیب از فیلد order (کامای جدا) یا checkboxها
        order_raw = (request.form.get('columns_order') or '').strip()
        if order_raw:
            keys = [k.strip() for k in order_raw.split(',') if k.strip()]
        else:
            keys = request.form.getlist('columns')
        # اجباری‌ها
        for req in ('actions', 'id', 'name'):
            if req not in keys:
                if req == 'actions':
                    keys.insert(0, req)
                elif req == 'id':
                    pos = keys.index('actions') + 1 if 'actions' in keys else 0
                    keys.insert(pos, req)
                else:
                    keys.append(req)
        # فیلتر معتبر
        valid = [k for k in keys if k in _CUSTOMER_COL_KEYS]
        # یکتا
        seen = set()
        final = []
        for k in valid:
            if k not in seen:
                final.append(k)
                seen.add(k)
        conn = get_db()
        set_setting(conn, CUSTOMER_LIST_COLUMNS_KEY, _json.dumps({'visible': final}, ensure_ascii=False))
        conn.commit()
        conn.close()
        return redirect('/settings/customers?cols_saved=1')



    @app.route('/settings/customers/edit/<int:customer_id>', methods=['GET', 'POST'])
    def edit_customer(customer_id):
        current_user = get_current_user()
        if not current_user or not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')

        conn = get_db()
        customer = conn.execute('SELECT * FROM customers WHERE id = ?', (customer_id,)).fetchone()
        if not customer:
            conn.close()
            return redirect('/settings/customers')

        if request.method == 'POST':
            name = request.form['name']
            phone = request.form['phone']
            national_id = request.form['national_id']
            economic_code = request.form['economic_code']
            province = request.form['province']
            county = request.form.get('county') or None
            city = request.form.get('city') or request.form.get('city_manual') or ''
            address = request.form['address']
            postal_code = request.form.get('postal_code', '')
            equipment_manager_name = request.form['equipment_manager_name']
            equipment_manager_mobile = request.form['equipment_manager_mobile']
            person_type = (request.form.get('person_type') or 'حقیقی').strip()
            if person_type not in ('حقیقی', 'حقوقی'):
                person_type = 'حقیقی'
            try:
                cols = {r[1] for r in conn.execute('PRAGMA table_info(customers)').fetchall()}
                if 'person_type' not in cols:
                    conn.execute("ALTER TABLE customers ADD COLUMN person_type TEXT DEFAULT 'حقیقی'")
            except Exception as e:
                logger.exception("Silent exception in routes/settings.py")
            conn.execute(
                '''UPDATE customers SET
                   name = ?, phone = ?, national_id = ?, economic_code = ?,
                   province = ?, county = ?, city = ?, address = ?, postal_code = ?,
                   equipment_manager_name = ?, equipment_manager_mobile = ?, person_type = ?
                   WHERE id = ?''',
                (name, phone, national_id, economic_code, province, county, city, address, postal_code,
                 equipment_manager_name, equipment_manager_mobile, person_type, customer_id)
            )
            conn.commit()
            conn.close()
            return redirect('/settings/customers')

        conn.close()
        return render_template(
            'edit_customer.html', customer=customer, provinces=(get_geo_provinces(conn) if get_geo_provinces else PROVINCES),
            active_page='settings', current_user=current_user, embed=_is_embed()
        )





    @app.route('/settings/customers/view/<int:customer_id>')
    def customer_profile(customer_id):
        """پروفایل مشتری — اطلاعات، آمار، دستگاه‌ها و درخواست‌ها."""
        current_user = get_current_user()
        if not current_user:
            return redirect('/login')
        if not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')

        conn = get_db()
        customer = conn.execute('SELECT * FROM customers WHERE id = ?', (customer_id,)).fetchone()
        if not customer:
            conn.close()
            return redirect('/settings/customers')

        name = customer['name'] or ''
        try:
            requests_rows = conn.execute(
                "SELECT id, reception_date, device_type, device_model, serial_number, status, "
                "service_type, request_category, problem_desc, province, city, county "
                "FROM requests WHERE customer_id = ? OR customer_name = ? ORDER BY id DESC LIMIT 100",
                (customer_id, name)
            ).fetchall()
        except Exception:
            requests_rows = conn.execute(
                "SELECT id, reception_date, device_type, serial_number, status, service_type, request_category "
                "FROM requests WHERE customer_name = ? ORDER BY id DESC LIMIT 100",
                (name,)
            ).fetchall()

        total = len(requests_rows)
        open_count = 0
        closed_words = ('پایان', 'تحویل', 'بایگانی', 'لغو')
        for r in requests_rows:
            st = (r['status'] or '')
            if not any(w in st for w in closed_words):
                open_count += 1

        devices_map = {}
        for r in requests_rows:
            keys = r.keys()
            serial = (r['serial_number'] or '').strip()
            dmodel = r['device_model'] if 'device_model' in keys else ''
            key = serial or ('%s|%s|%s' % (r['device_type'] or '', dmodel or '', r['id']))
            if key in devices_map:
                continue
            photo = None
            try:
                dt = r['device_type'] or ''
                if dt and dmodel:
                    drow = conn.execute(
                        'SELECT photo_filename FROM devices WHERE device_type=? AND model=? LIMIT 1',
                        (dt, dmodel)
                    ).fetchone()
                elif dt:
                    drow = conn.execute(
                        'SELECT photo_filename FROM devices WHERE device_type=? LIMIT 1', (dt,)
                    ).fetchone()
                else:
                    drow = None
                if drow and drow['photo_filename']:
                    photo = drow['photo_filename']
            except Exception as e:
                logger.exception("Silent exception in routes/settings.py")
            devices_map[key] = {
                'device_type': r['device_type'],
                'device_model': dmodel,
                'serial_number': serial,
                'last_status': r['status'],
                'last_date': r['reception_date'] if 'reception_date' in keys else '',
                'photo_filename': photo,
                'last_request_id': r['id'],
                'service_type': r['service_type'] if 'service_type' in keys else '',
            }
        customer_devices = list(devices_map.values())

        month_counts = {}
        for r in requests_rows:
            rd = (r['reception_date'] or '')[:7]
            if len(rd) >= 7 and rd[4:5] == '/':
                month_counts[rd] = month_counts.get(rd, 0) + 1
        chart_labels = sorted(month_counts.keys())[-6:]
        chart_values = [month_counts[k] for k in chart_labels]

        warranties = []
        try:
            warranties = conn.execute(
                'SELECT * FROM warranty_cards WHERE customer_name = ? ORDER BY id DESC LIMIT 50',
                (name,)
            ).fetchall()
        except Exception:
            warranties = []

        conn.close()
        return render_template(
            'customer_profile.html',
            customer=customer,
            requests_rows=requests_rows,
            customer_devices=customer_devices,
            warranties=warranties,
            stats={
                'total': total,
                'open': open_count,
                'closed': max(0, total - open_count),
                'devices': len(customer_devices),
            },
            chart_labels=chart_labels,
            chart_values=chart_values,
            active_page='settings',
            current_user=current_user,
        )



    @app.route('/settings/customers/delete/<int:customer_id>', methods=['POST'])
    def delete_customer(customer_id):
        """آرشیو نرم — حذف فیزیکی انجام نمی‌شود."""
        current_user = get_current_user()
        if not current_user or not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')

        conn = get_db()
        try:
            cols = {r[1] for r in conn.execute('PRAGMA table_info(customers)').fetchall()}
            if 'is_archived' not in cols:
                conn.execute('ALTER TABLE customers ADD COLUMN is_archived INTEGER DEFAULT 0')
            if 'archive_date' not in cols:
                conn.execute('ALTER TABLE customers ADD COLUMN archive_date TEXT')
            try:
                import jdatetime
                td = jdatetime.date.today()
                ad = f"{td.year}/{str(td.month).zfill(2)}/{str(td.day).zfill(2)}"
            except Exception:
                ad = None
            conn.execute(
                'UPDATE customers SET is_archived=1, archive_date=COALESCE(?, archive_date) WHERE id=?',
                (ad, customer_id),
            )
            try:
                write_audit(conn, current_user.get('full_name'), 'آرشیو طرف حساب', 'customer', customer_id, '')
            except Exception as e:
                logger.exception("Silent exception in routes/settings.py")
            conn.commit()
        except Exception as e:
            logger.exception("Silent exception in routes/settings.py")
        conn.close()
        return redirect('/settings/customers')





    @app.route('/settings/customers/export')
    def export_customers():
        current_user = get_current_user()
        if not current_user or not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')

        conn = get_db()
        customers_list = conn.execute('SELECT * FROM customers WHERE IFNULL(is_archived,0)=0 ORDER BY id').fetchall()
        conn.close()

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'طرف حساب ها'

        headers = [
            'کد طرف حساب', 'نوع', 'نام مشتری', 'استان', 'شهرستان', 'شهر', 'آدرس', 'کد پستی',
            'کد ملی / شناسه ملی', 'شناسه اقتصادی', 'تلفن ثابت', 'نام مسئول تجهیزات', 'تلفن همراه'
        ]
        sheet.append(headers)

        for c in customers_list:
            keys = c.keys()
            sheet.append([
                c['id'],
                c['person_type'] if 'person_type' in keys and c['person_type'] else 'حقیقی',
                c['name'], c['province'],
                c['county'] if 'county' in keys else '',
                c['city'], c['address'], c['postal_code'],
                c['national_id'], c['economic_code'], c['phone'],
                c['equipment_manager_name'], c['equipment_manager_mobile']
            ])

        for col in sheet.columns:
            max_length = max((len(str(cell.value)) if cell.value else 0) for cell in col)
            sheet.column_dimensions[col[0].column_letter].width = max_length + 4

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='customers.xlsx'
        )





    @app.route('/settings/customers/import', methods=['POST'])
    def import_customers():
        current_user = get_current_user()
        if not current_user or not (_is_reception(current_user) or user_has_access(current_user, 'parties')):
            return redirect('/')

        file = request.files.get('excel_file')
        import_error = None
        imported_count = 0

        if not file or file.filename == '':
            import_error = 'هیچ فایلی انتخاب نشده است.'
        elif not file.filename.lower().endswith(('.xlsx', '.xls')):
            import_error = 'فقط فایل اکسل (xlsx یا xls) مجاز است.'
        else:
            try:
                workbook = openpyxl.load_workbook(file, data_only=True)
                sheet = workbook.active
                conn = get_db()

                rows = list(sheet.iter_rows(values_only=True))
                data_rows = rows[1:] if rows else []

                # ترتیب ستون‌ها (نسخه 2.10):
                # کد طرف حساب، نوع (حقیقی/حقوقی)، نام مشتری، استان، شهرستان، شهر، آدرس، کد پستی،
                # کد ملی/شناسه ملی، شناسه اقتصادی، تلفن ثابت، نام مسئول تجهیزات، تلفن همراه
                # سازگاری با فایل قدیمی (بدون ستون نوع): اگر ستون۲ نام باشد، نوع=حقیقی فرض می‌شود.
                for row in data_rows:
                    if not row:
                        continue
                    values = list(row) + [None] * 16
                    # تشخیص فرمت: اگر values[1] یکی از حقیقی/حقوقی باشد، فرمت جدید است
                    v1 = (str(values[1]).strip() if values[1] is not None else '')
                    if v1 in ('حقیقی', 'حقوقی'):
                        person_type = v1
                        name = values[2]
                        province = values[3]
                        county = values[4]
                        city = values[5]
                        address = values[6]
                        postal_code = values[7]
                        national_id = values[8]
                        economic_code = values[9]
                        phone = values[10]
                        equipment_manager_name = values[11]
                        equipment_manager_mobile = values[12]
                    else:
                        # فرمت قدیمی
                        if values[1] is None:
                            continue
                        person_type = 'حقیقی'
                        name = values[1]
                        province = values[2]
                        county = values[3]
                        city = values[4]
                        address = values[5]
                        national_id = values[6]
                        economic_code = values[7]
                        phone = values[8]
                        equipment_manager_name = values[9]
                        equipment_manager_mobile = values[10]
                        postal_code = values[11]
                    if name is None or str(name).strip() == '':
                        continue
                    try:
                        cols = {r[1] for r in conn.execute('PRAGMA table_info(customers)').fetchall()}
                        if 'person_type' not in cols:
                            conn.execute("ALTER TABLE customers ADD COLUMN person_type TEXT DEFAULT 'حقیقی'")
                    except Exception as e:
                        logger.exception("Silent exception in routes/settings.py")
                    conn.execute(
                        '''INSERT INTO customers
                           (name, phone, national_id, economic_code, province, county, city, address, postal_code,
                            equipment_manager_name, equipment_manager_mobile, person_type)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                        (str(name), str(phone) if phone else None, str(national_id) if national_id else None,
                         str(economic_code) if economic_code else None, str(province) if province else None,
                         str(county) if county else None,
                         str(city) if city else None, str(address) if address else None,
                         str(postal_code) if postal_code else None,
                         str(equipment_manager_name) if equipment_manager_name else None,
                         str(equipment_manager_mobile) if equipment_manager_mobile else None,
                         person_type)
                    )
                    imported_count += 1

                conn.commit()
                conn.close()
            except Exception as e:
                import_error = f'خطا در پردازش فایل: {e}'

        conn = get_db()
        customers_list = [dict(r) for r in conn.execute('SELECT * FROM customers WHERE IFNULL(is_archived,0)=0 ORDER BY id DESC').fetchall()]
        provinces_list = (get_geo_provinces(conn) if get_geo_provinces else PROVINCES)
        visible_cols = _customer_columns_for_template(conn)
        can_cols = _can_edit_customer_columns(current_user)
        conn.close()
        return render_template(
            'settings_customers.html',
            customers=customers_list,
            provinces=provinces_list,
            active_page='settings',
            current_user=current_user,
            import_error=import_error,
            imported_count=imported_count,
            customer_columns=visible_cols,
            can_edit_columns=can_cols,
            all_column_defs=CUSTOMER_LIST_COLUMN_DEFS,
        )




